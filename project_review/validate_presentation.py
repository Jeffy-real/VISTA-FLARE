import os
import re
import sys
from pathlib import Path
import pptx
from pptx import Presentation

PPTX_PATH = Path("project_review/Vista_Flare_Deep_Learning_Project_Review_2.pptx")

TEMPLATE_LEAKS = [
    "Write 10-12 lines",
    "Write 3-4 points",
    "Mention 3-4 key points",
    "Write 3-4 objectives",
    "Refer minimum 15 latest research papers",
    "Which model you are going to use",
    "Reason for choosing your model",
    "Mention the dataset name that you are going to use",
    "Draw the architecture diagram",
    "You can add the slides",
    "Screenshots like bar-graphs/charts",
    "Dr. xxxx",
    "S. Kumar, P. Tiwari and M. Zymbler",
    "Kindly follow the same reference format"
]

SECRET_PATTERNS = [
    r"sk-[a-zA-Z0-9]{20,}",
    r"eyJ[a-zA-Z0-9_-]{20,}",
    r"postgres://",
    r"password\s*=\s*['\"][^'\"]+['\"]",
    r"api_key\s*=\s*['\"][^'\"]+['\"]"
]

def validate():
    print(f"=== VALIDATING PPTX: {PPTX_PATH} ===")
    assert PPTX_PATH.exists(), f"File does not exist at {PPTX_PATH}"
    print(f"[PASS] File exists. Size: {PPTX_PATH.stat().st_size:,} bytes")

    prs = Presentation(str(PPTX_PATH))
    total_slides = len(prs.slides)
    print(f"[PASS] Successfully opened presentation. Total slides: {total_slides}")
    assert total_slides >= 19, f"Expected at least 19 slides, found {total_slides}"

    leaks_found = []
    secrets_found = []
    slide_summary = []

    for idx, slide in enumerate(prs.slides):
        title = "No Title"
        text_corpus = []
        table_count = 0
        image_count = 0

        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text.strip()
                if txt:
                    text_corpus.append(txt)
                    if not title or title == "No Title":
                        lines = [line.strip() for line in txt.split("\n") if line.strip()]
                        if lines:
                            title = lines[0]
            if shape.has_table:
                table_count += 1
                tbl = shape.table
                for row in tbl.rows:
                    row_txt = " | ".join(c.text.strip() for c in row.cells)
                    text_corpus.append(row_txt)
            if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                image_count += 1

        full_slide_text = "\n".join(text_corpus)

        # Check for unresolved template text
        for pat in TEMPLATE_LEAKS:
            if pat.lower() in full_slide_text.lower():
                leaks_found.append((idx + 1, pat))

        # Check for secret leaks
        for pat in SECRET_PATTERNS:
            if re.search(pat, full_slide_text):
                secrets_found.append((idx + 1, pat))

        slide_summary.append({
            "slide": idx + 1,
            "title": title[:50],
            "text_length": len(full_slide_text),
            "tables": table_count,
            "images": image_count
        })

    print("\n--- SLIDE INVENTORY ---")
    for s in slide_summary:
        print(f"Slide {s['slide']:02d}: Title='{s['title']}' | TextLen={s['text_length']} | Tables={s['tables']} | Images={s['images']}")

    print("\n--- TEMPLATE PLACEHOLDER CHECK ---")
    if leaks_found:
        print(f"[FAIL] Found {len(leaks_found)} unresolved template placeholder fragments:")
        for slide_no, pat in leaks_found:
            print(f"   Slide {slide_no}: contains '{pat}'")
    else:
        print("[PASS] Zero unresolved template placeholder text found!")

    print("\n--- SECRET & CREDENTIAL LEAK CHECK ---")
    if secrets_found:
        print(f"[FAIL] Found potential credentials on slides: {secrets_found}")
    else:
        print("[PASS] Zero secrets or credentials detected in any slide text!")

    print("\n--- SECTION SEQUENCE VERIFICATION ---")
    expected_sections = [
        "School of Computer Science and Engineering",
        "Outline",
        "ABSTRACT",
        "INTRODUCTION",
        "PROBLEM STATEMENT",
        "OBJECTIVES",
        "LITERATURE REVIEW",
        "Software and Hardware Requirements",
        "PROPOSED SOLUTION",
        "DATASET DETAILS",
        "METHODOLOGY",
        "MODEL EXPLANATION",
        "RESULTS",
        "CONCLUSION",
        "FUTURE ENHANCEMENTS",
        "REFERENCES"
    ]
    all_titles = " ".join(s['title'] for s in slide_summary)
    for sec in expected_sections:
        found = any(sec.lower() in s['title'].lower() for s in slide_summary)
        print(f"   Section '{sec}': {'FOUND [PASS]' if found else 'MISSING [FAIL]'}")

    assert not leaks_found, "Template placeholders remain"
    assert not secrets_found, "Secrets detected"
    print("\n=== ALL VALIDATION ASSERTIONS PASSED! ===")

if __name__ == "__main__":
    validate()

